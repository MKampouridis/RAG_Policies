"""Readable text from a .docx, for both the local ingester and the crawler.

Extracted from ingest_local.py (2026-09-08) so the crawler can parse a docx it
fetched over HTTP without duplicating the parser. Same reason it was written
carefully in the first place: Essex publishes some policy as Word, and the
substance of those documents sits in tables as much as in prose.

Takes a path OR a file-like object, because ingest_local has a file on disk and
the crawler has response bytes.

NOT .doc: python-docx reads the Office Open XML format only, and the legacy
binary .doc is a different format entirely. Callers should check the extension
and skip rather than hand one over.
"""

import docx


def _iter_block_text(document) -> list[str]:
    """Paragraphs and tables in document order.

    python-docx exposes .paragraphs and .tables as separate flat lists, which
    loses their interleaving - and these documents' substance is largely in
    appendix tables sitting between explanatory paragraphs. Walking the body
    XML keeps reading order, so a chunk boundary can't land between a table
    and the sentence that introduces it.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    out = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            out.append(Paragraph(child, document).text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                # de-duplicate horizontally merged cells, which python-docx
                # reports once per underlying grid column
                deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                line = " | ".join(c for c in deduped if c)
                if line:
                    out.append(line)
    return out


def extract_docx_text(source) -> str:
    """Readable text from a .docx (path or file-like), minus Word's
    table-of-contents plumbing.

    TOC entries survive as literal field text ("Introduction PAGEREF
    _Toc234314569 \\h 2"). They are pure noise for retrieval - a list of
    headings the body already contains, carrying page numbers that mean
    nothing once chunked - and they would otherwise be the document's most
    heading-dense chunk, which is exactly the shape that wins on identity
    queries while answering nothing.
    """
    document = docx.Document(source if hasattr(source, "read") else str(source))
    lines = []
    for raw in _iter_block_text(document):
        line = raw.strip()
        if not line:
            continue
        if "PAGEREF" in line or line.startswith("TOC \\"):
            continue
        lines.append(line)
    return "\n".join(lines)
